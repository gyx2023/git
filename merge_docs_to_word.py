#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import markdown
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 文档目录
DOCS_DIR = "/Users/liutong/farben/ruoyi-flowable-plus-0.8.X/docs"

# 文档顺序
DOCS_ORDER = [
    "README.md",
    "项目概述.md",
    "系统架构文档.md",
    "核心功能模块文档.md",
    "API接口文档.md",
    "部署和运行指南.md",
    "用户使用手册.md",
    "配置文件说明.md"
]

def read_markdown_file(file_path):
    """读取Markdown文件内容"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except Exception as e:
        print(f"读取文件 {file_path} 失败: {e}")
        return ""

def add_title_to_doc(doc, title, level=1):
    """添加标题到Word文档"""
    if level == 1:
        # 一级标题居中
        para = doc.add_heading(title, level=level)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_heading(title, level=level)

def parse_markdown_to_doc(md_content, doc, skip_title=False):
    """解析Markdown内容并添加到Word文档"""
    lines = md_content.split('\n')
    in_code_block = False
    code_content = []
    table_mode = False
    table_lines = []
    current_heading = None

    for line in lines:
        # 跳过空行
        if not line.strip():
            if code_content:
                code_content.append('')
            elif not in_code_block and not table_mode:
                doc.add_paragraph()
            continue

        # 处理代码块
        if line.strip().startswith('```'):
            if in_code_block:
                # 结束代码块
                in_code_block = False
                if code_content:
                    code_text = '\n'.join(code_content)
                    # 添加代码块（作为普通文本）
                    para = doc.add_paragraph()
                    run = para.add_run(code_text)
                    run.font.name = 'Consolas'
                    para.paragraph_format.left_indent = Inches(0.5)
                    code_content = []
            else:
                # 开始代码块
                in_code_block = True
                # 获取代码语言（如果有）
                code_language = line.strip()[3:].strip()
                if code_language:
                    doc.add_paragraph(f"代码语言: {code_language}")
            continue

        # 收集代码块内容
        if in_code_block:
            code_content.append(line)
            continue

        # 处理表格
        if '|' in line and not line.strip().startswith('#') and not line.strip().startswith('- [') and not line.strip().startswith('* '):
            if not table_mode:
                table_mode = True
            table_lines.append(line)
            continue
        elif table_mode:
            # 处理表格
            process_table(table_lines, doc)
            table_mode = False
            table_lines = []

        # 处理标题
        if line.startswith('#'):
            level = len(line.split(' ')[0])
            title_text = line[level:].strip()
            if skip_title and level == 1:
                current_heading = title_text
                continue
            add_title_to_doc(doc, title_text, level)
            continue

        # 处理列表
        if line.strip().startswith(('- [', '* [', '- ', '* ', '1. ', '2. ', '3. ')):
            # 简单处理为段落，前面添加缩进
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.5)
            para.add_run(line.strip())
            continue

        # 处理普通段落
        doc.add_paragraph(line.strip())

    # 处理剩余的表格
    if table_mode and table_lines:
        process_table(table_lines, doc)

def process_table(table_lines, doc):
    """处理表格内容"""
    if not table_lines:
        return

    # 简单地将表格内容作为文本添加
    # 由于docx表格处理较复杂，这里简化处理
    doc.add_paragraph("表格内容:")
    for line in table_lines:
        para = doc.add_paragraph()
        para.add_run(line)

def generate_word_with_toc(docs_dir, docs_order, output_file):
    """生成带目录的Word文档"""
    doc = Document()

    # 设置文档标题
    title_para = doc.add_heading('zqzt_backend 项目文档', level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_paragraph()

    # 添加目录
    doc.add_heading('目录', level=1)
    toc_items = []

    # 收集所有标题
    for doc_file in docs_order:
        file_path = os.path.join(docs_dir, doc_file)
        if not os.path.exists(file_path):
            print(f"文件不存在: {file_path}")
            continue

        md_content = read_markdown_file(file_path)
        lines = md_content.split('\n')

        # 获取文档主标题（一级标题）
        main_title = None
        for line in lines:
            if line.startswith('# '):
                main_title = line[2:].strip()
                break

        if main_title:
            toc_items.append((main_title, 1))

            # 收集二级和三级标题
            for line in lines:
                if line.startswith('## '):
                    toc_items.append((line[3:].strip(), 2))
                elif line.startswith('### '):
                    toc_items.append((line[4:].strip(), 3))

    # 生成目录内容
    for title, level in toc_items:
        para = doc.add_paragraph()
        if level == 1:
            para.add_run(title).bold = True
        elif level == 2:
            para.paragraph_format.left_indent = Inches(0.5)
            para.add_run(title)
        elif level == 3:
            para.paragraph_format.left_indent = Inches(1.0)
            para.add_run(title)

    doc.add_page_break()

    # 添加文档内容
    for doc_file in docs_order:
        file_path = os.path.join(docs_dir, doc_file)
        if not os.path.exists(file_path):
            print(f"文件不存在: {file_path}")
            continue

        print(f"处理文档: {doc_file}")
        md_content = read_markdown_file(file_path)

        # 为README.md特殊处理，跳过它的一级标题
        skip_title = doc_file == "README.md"
        parse_markdown_to_doc(md_content, doc, skip_title=skip_title)
        doc.add_page_break()

    # 保存文档
    doc.save(output_file)
    print(f"Word文档已生成: {output_file}")

if __name__ == "__main__":
    output_file = "/Users/liutong/farben/ruoyi-flowable-plus-0.8.X/zqzt_backend_项目文档.docx"
    generate_word_with_toc(DOCS_DIR, DOCS_ORDER, output_file)
    print("文档合并完成！")
